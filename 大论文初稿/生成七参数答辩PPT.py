#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""生成七参数 MLP-CNN 流热场重建答辩 PPT。"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "大论文初稿" / "ppt模版" / "毕业答辩PPT_final.pptx"
OUT = ROOT / "大论文初稿" / "七参数MLP-CNN流热场重建答辩PPT.pptx"

INK = RGBColor(28, 35, 43)
MUTED = RGBColor(91, 104, 118)
ACCENT = RGBColor(26, 115, 232)
GREEN = RGBColor(20, 125, 85)
ORANGE = RGBColor(214, 111, 37)
LIGHT = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, left, top, width, height, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, kicker, title):
    add_text(slide, kicker, Inches(0.55), Inches(0.25), Inches(2.0), Inches(0.25), 10, True, ACCENT)
    add_text(slide, title, Inches(0.55), Inches(0.55), Inches(11.0), Inches(0.55), 25, True, INK)
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(1.16), Inches(11.6), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 226, 233)
    line.line.fill.background()


def add_footer(slide, page):
    add_text(slide, f"{page:02d}", Inches(11.8), Inches(6.85), Inches(0.55), Inches(0.2), 9, True, MUTED, PP_ALIGN.RIGHT)


def add_bullets(slide, items, left, top, width, height, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.space_after = Pt(6)
    return box


def add_card(slide, title, body, left, top, width, height, color=LIGHT):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(218, 225, 233)
    add_text(slide, title, left + Inches(0.18), top + Inches(0.12), width - Inches(0.36), Inches(0.25), 13, True, INK)
    add_text(slide, body, left + Inches(0.18), top + Inches(0.45), width - Inches(0.36), height - Inches(0.55), 12, False, MUTED)


def add_metric(slide, value, label, left, top, width=Inches(2.2), accent=ACCENT):
    add_text(slide, value, left, top, width, Inches(0.45), 24, True, accent, PP_ALIGN.CENTER)
    add_text(slide, label, left, top + Inches(0.42), width, Inches(0.25), 10, False, MUTED, PP_ALIGN.CENTER)


def add_table(slide, headers, rows, left, top, width, height, font_size=10):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height).table
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(230, 236, 244)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(font_size)
                paragraph.font.color.rgb = INK
    return table


def add_picture_if_exists(slide, path, left, top, width=None, height=None):
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    else:
        add_card(slide, "图片缺失", str(path.relative_to(ROOT)), left, top, width or Inches(4), height or Inches(2))


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def main() -> None:
    prs = Presentation(str(TEMPLATE)) if TEMPLATE.exists() else Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    while len(prs.slides) > 0:
        xml_slides = prs.slides._sldIdLst
        rel_id = xml_slides[0].rId
        prs.part.drop_rel(rel_id)
        xml_slides.remove(xml_slides[0])

    slides = []
    for _ in range(23):
        slides.append(blank_slide(prs))

    # 1
    s = slides[0]
    add_text(s, "毕业论文答辩", Inches(0.7), Inches(0.55), Inches(3), Inches(0.35), 15, True, ACCENT)
    add_text(s, "基于 MLP-CNN 流热场重建的\n翼型柱阵列散热器性能预测与优化", Inches(0.7), Inches(1.35), Inches(10.8), Inches(1.3), 34, True, INK)
    add_text(s, "七参数参数化建模 · PVT 流热场重建 · DE 优化 · CFD 复算 · 三维迁移验证", Inches(0.75), Inches(3.05), Inches(11), Inches(0.45), 17, False, MUTED)
    add_metric(s, "1000", "七参数 CFD 样本", Inches(0.8), Inches(4.45))
    add_metric(s, "3×96×320", "PVT 场数据维度", Inches(3.25), Inches(4.45), accent=GREEN)
    add_metric(s, "0.971888", "eta_R2", Inches(5.9), Inches(4.45), accent=ORANGE)
    add_metric(s, "1.056732", "三维 eta3D", Inches(8.25), Inches(4.45), accent=ACCENT)
    add_footer(s, 1)

    content = [
        ("研究背景", "七参数结构优化需要比逐点 CFD 更快的预测方法", ["翼型柱阵列同时影响换热增强和流动阻力。", "七维参数空间直接 CFD 遍历成本高。", "只预测 Nu/f 缺少流热场解释能力。"]),
        ("问题定义", "目标是同时预测性能指标和 p/U/T 空间流热场", ["输入：Ta、Twa、Tb、Ts、Tt、Tad、theta。", "输出：p/U/T 场、Nu、f，eta 由 Nu/f 计算。", "PVT 中 V 表示速度大小 U=sqrt(u²+v²)。"]),
        ("技术路线", "从参数化建模到三维迁移形成完整证据链", ["七参数建模 -> 二维 CFD -> PVT 数据集。", "MLP-CNN / Conditional U-Net 重建流热场。", "DE 搜索 eta 最大结构，二维复算后迁移三维。"]),
        ("七参数建模", "几何参数控制翼型形状、阵列间距和整体姿态", ["Ta/Twa/Tb 描述翼型几何。", "Ts/Tt/Tad 控制阵列间距与错排。", "theta 控制整体旋转角。"]),
        ("二维 CFD 设置", "二维复算用于验证代理优化结论", ["入口给定流速和温度，出口压力边界。", "固体/流体区域耦合传热。", "输出 Nu、f、eta 以及速度、压力、温度场。"]),
        ("PVT 数据集", "1000 组样本统一为 3×96×320 高分辨率场", ["数据文件：field_reconstruction_dataset_320x96_full_pUt.npz。", "训练/验证/测试划分：700/150/150。", "CFD 场是监督标签，不是模型输入。"]),
        ("不是几何掩码", "最终主线不依赖几何掩码图像输入", ["几何掩码容易被理解为参数曲线重绘。", "本文模型输入仍是七参数和空间条件。", "CNN 的任务是解码 p/U/T 空间场。"]),
        ("网络结构", "MLP 编码参数，CNN/U-Net 解码空间场", ["MLP Encoder：七参数 -> 潜在特征。", "CNN Decoder：潜在特征 + 坐标/区域 -> p/U/T。", "性能分支输出 Nu/f，eta 保持公式一致。"]),
        ("损失函数", "多任务损失同时约束场、指标和物理一致性", ["PVT 场重建损失。", "Nu/f 性能预测损失。", "eta 一致性损失和梯度一致性损失。"]),
    ]
    for i, (kicker, title, bullets) in enumerate(content, start=2):
        s = slides[i - 1]
        add_title(s, kicker, title)
        add_bullets(s, bullets, Inches(0.85), Inches(1.65), Inches(5.1), Inches(3.6), 18)
        add_card(s, "证据路径", "论文正文与 reports/论文证据链总表.md 中列出脚本、JSON、CSV 和图片路径。", Inches(6.55), Inches(1.75), Inches(5.4), Inches(1.2))
        add_footer(s, i)

    # 11 主模型
    s = slides[10]
    add_title(s, "主模型结果", "PVT U-Net 同时保持标量预测和场重建精度")
    add_metric(s, "0.969309", "Nu_R2", Inches(0.75), Inches(1.7))
    add_metric(s, "0.978038", "f_R2", Inches(3.15), Inches(1.7), accent=GREEN)
    add_metric(s, "0.971888", "eta_R2", Inches(5.55), Inches(1.7), accent=ORANGE)
    add_metric(s, "0.439960 K", "T_MAE", Inches(8.15), Inches(1.7), accent=ACCENT)
    add_picture_if_exists(s, ROOT / "七参数_MLP_CNN流场重建/results/unet_pUt_320x96_full_e50/prediction_scatter.png", Inches(1.0), Inches(3.1), width=Inches(10.8))
    add_footer(s, 11)

    # 12 PVT 云图
    s = slides[11]
    add_title(s, "PVT 云图对比", "预测场可用于解释压力、速度和温度分布趋势")
    add_picture_if_exists(s, ROOT / "七参数_MLP_CNN流场重建/results/pvt_comsol_style_exports_pUt_full_e50_case20040/case_20040_T_comsol_sample_style.png", Inches(0.7), Inches(1.55), width=Inches(3.8))
    add_picture_if_exists(s, ROOT / "七参数_MLP_CNN流场重建/results/pvt_comsol_style_exports_pUt_full_e50_case20040/case_20040_T_pred_style.png", Inches(4.75), Inches(1.55), width=Inches(3.8))
    add_picture_if_exists(s, ROOT / "七参数_MLP_CNN流场重建/results/pvt_comsol_style_exports_pUt_full_e50_case20040/case_20040_T_error_style.png", Inches(8.8), Inches(1.55), width=Inches(3.8))
    add_footer(s, 12)

    # 13 K fold
    s = slides[12]
    add_title(s, "K 折验证", "三折结果说明模型对数据划分较稳定")
    add_table(s, ["指标", "均值", "标准差"], [
        ["Nu_R2", "0.964261", "0.003353"],
        ["f_R2", "0.969705", "0.004482"],
        ["eta_R2", "0.968501", "0.000825"],
        ["T_MAE", "0.614454", "0.049592"],
        ["field_MAE", "0.649893", "0.028645"],
    ], Inches(1.0), Inches(1.65), Inches(6.3), Inches(2.2), 12)
    add_card(s, "结论", "eta_R2 标准差仅 0.000825，说明综合性能预测对数据划分不敏感；K 折不替代 50 轮全量主模型。", Inches(8.0), Inches(1.75), Inches(4.3), Inches(1.6))
    add_footer(s, 13)

    # 14 ablation
    s = slides[13]
    add_title(s, "消融实验", "条件输入是场重建精度提升的关键")
    add_table(s, ["组别", "eta_R2", "field_MAE", "T_MAE", "结论"], [
        ["A0", "0.969789", "-", "-", "仅标量"],
        ["A1", "0.974212", "3.708562", "3.316894", "可重建但误差大"],
        ["A2", "0.973342", "0.616973", "0.637812", "条件输入有效"],
        ["A3", "0.973451", "0.608167", "0.509703", "物理约束改善局部"],
        ["A4", "0.972475", "0.571821", "0.469610", "CBAM 改善场误差"],
    ], Inches(0.65), Inches(1.45), Inches(12.0), Inches(2.7), 10)
    add_footer(s, 14)

    # 15 sensitivity
    s = slides[14]
    add_title(s, "参数敏感性", "Ts 是 eta 的最显著敏感参数")
    add_picture_if_exists(s, ROOT / "七参数_MLP_CNN流场重建/analysis_results/sensitivity/sensitivity_bar_eta.png", Inches(0.8), Inches(1.55), width=Inches(6.2))
    add_table(s, ["排序", "参数", "证据"], [
        ["1", "Ts", "RF permutation=1.361628"],
        ["2", "theta", "RF permutation=0.231464"],
        ["3", "Tb", "RF permutation=0.147826"],
    ], Inches(7.4), Inches(1.75), Inches(4.8), Inches(1.65), 11)
    add_footer(s, 15)

    # 16 grid
    s = slides[15]
    add_title(s, "网格无关性", "二维和三维三档网格已补算，但当前未通过阈值")
    add_bullets(s, ["二维 G4/G5：baseline Nu 21.070%、f 14.329%；DE Nu 16.006%、f 11.730%。", "三维 Tmax 已接近收敛，但压降差异仍为 13.385% / 8.063%。", "诊断：域级局部尺寸细化后二维仍未过，后续需壁面边界层和尾迹局部控制。"], Inches(0.85), Inches(1.6), Inches(7.35), Inches(3.0), 16)
    add_card(s, "写法边界", "不能写成已完成网格无关性；当前只能报告真实补算未通过，并以正式 CFD 复算和三维迁移作为性能证据。", Inches(8.45), Inches(1.75), Inches(3.85), Inches(1.55), RGBColor(255, 246, 230))
    add_footer(s, 16)

    # 17-23
    tail = [
        ("DE 优化", "代理模型搜索 eta 最大的七参数组合", ["最优：Ta=0, Twa=0.6, Tb=0.15, Ts=1.5。", "Tt=0.960919, Tad=1.0, theta=-2.224429。", "代理预测 eta_pred=1.433572。"]),
        ("二维 CFD 复算", "最终优化结论以 COMSOL 复算为准", ["Nu_CFD=46.022427。", "f_CFD=0.106658。", "eta_CFD=1.447895，高于旧 baseline eta=1.420111。"]),
        ("三维迁移验证", "二维最优结构在芯片级三维模型中仍有正收益", ["Tavg 从 379.201 K 降至 375.663 K。", "Tmax 从 380.457 K 降至 376.970 K。", "eta3D=1.056732；这是迁移验证，不是三维全局寻优结果。"]),
        ("结果讨论", "模型有效，但不等于高保真 CFD 求解器", ["PVT 场能解释主要空间趋势。", "翼型边界、尾迹和压力突变仍有平滑误差。", "优化结论必须以 CFD 复算为准。"]),
        ("创新点", "本文贡献在数据链、场重建和优化验证的一体化", ["七参数到 PVT 场的 MLP-CNN/U-Net 代理。", "性能预测与流热场监督联合学习。", "DE + 二维复算 + 三维迁移的证据链。"]),
        ("结论", "优化结构在二维和三维验证中均表现出正向收益", ["PVT 主模型 eta_R2=0.971888。", "二维 CFD eta=1.447895。", "三维 eta3D=1.056732。"]),
        ("答辩问答", "核心问题要围绕证据链回答", ["为什么不用几何掩码？最终输入不是云图，云图是监督标签。", "网格为什么没写通过？真实补算未达阈值，只能写网格敏感性风险。", "为什么不替代 CFD？局部高梯度仍平滑，最终结论靠复算。"]),
    ]
    for offset, (kicker, title, bullets) in enumerate(tail, start=17):
        s = slides[offset - 1]
        add_title(s, kicker, title)
        add_bullets(s, bullets, Inches(0.85), Inches(1.65), Inches(6.2), Inches(3.8), 18)
        if offset == 19:
            add_picture_if_exists(s, ROOT / "comsol_3d_airfoil_radiator/generated_chip_airfoil_heat_sink/chip_mlp_cnn_de_multiloss_e10/exports/chip_mlp_cnn_de_multiloss_e10_temperature.png", Inches(7.4), Inches(1.55), width=Inches(4.7))
        else:
            add_card(s, "证据", "详见论文附录A与 reports/论文证据链总表.md。", Inches(7.5), Inches(1.8), Inches(4.2), Inches(1.2))
        add_footer(s, offset)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
