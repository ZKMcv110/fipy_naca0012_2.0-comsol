#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""生成当前论文主线答辩 PPT。

只读取新目录结果和三维迁移验证结果，不修改旧模板。
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT.parent
OUTPUT = ROOT / "MLP_CNN流热场重建_论文答辩PPT.pptx"

TRAIN_DIR = ROOT / "results" / "mlp_cnn_field_multiloss_e10"
OPT_DIR = ROOT / "optimization_results" / "mlp_cnn_de_multiloss_e10"
VAL_DIR = ROOT / "validation_results" / "mlp_cnn_de_multiloss_e10"
SENS_DIR = ROOT / "analysis_results" / "sensitivity"
KFOLD_DIR = ROOT / "results" / "mlp_cnn_kfold_3_e5"
SUMMARY_3D = (
    PROJECT_ROOT
    / "comsol_3d_airfoil_radiator"
    / "generated_chip_airfoil_heat_sink"
    / "results"
    / "seven_param_3d_summary.csv"
)
CASE_3D = "chip_mlp_cnn_de_multiloss_e10"
IMG_3D = (
    PROJECT_ROOT
    / "comsol_3d_airfoil_radiator"
    / "generated_chip_airfoil_heat_sink"
    / CASE_3D
    / "exports"
)

NAVY = RGBColor(24, 49, 83)
TEAL = RGBColor(0, 128, 128)
GRAY = RGBColor(90, 98, 108)
LIGHT = RGBColor(246, 248, 250)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(30, 30, 30)


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def load_3d_rows() -> tuple[dict, dict]:
    baseline, optimized = {}, {}
    with SUMMARY_3D.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row.get("case") == "baseline":
                baseline = row
            if row.get("case") == CASE_3D:
                optimized = row
    return baseline, optimized


def fnum(value, digits: int = 6) -> str:
    return f"{float(value):.{digits}f}"


def pct_change(new, old) -> str:
    return f"{(float(new) / float(old) - 1.0) * 100.0:+.3f}%"


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Cm(1.0), Cm(0.45), Cm(31.8), Cm(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Cm(1.05), Cm(1.45), Cm(30.8), Cm(0.55))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Microsoft YaHei"
        sp.font.size = Pt(10)
        sp.font.color.rgb = GRAY


def add_footer(slide, idx: int) -> None:
    box = slide.shapes.add_textbox(Cm(28.5), Cm(18.25), Cm(4.0), Cm(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = f"{idx:02d}"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.RIGHT


def new_slide(prs: Presentation, title: str, subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title, subtitle)
    add_footer(slide, len(prs.slides))
    return slide


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 15) -> None:
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = BLACK
        p.space_after = Pt(6)


def add_panel(slide, x: float, y: float, w: float, h: float, title: str, body: list[str]) -> None:
    shape = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(215, 222, 230)
    head = slide.shapes.add_textbox(Cm(x + 0.35), Cm(y + 0.25), Cm(w - 0.7), Cm(0.5))
    p = head.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    add_bullets(slide, body, x + 0.35, y + 0.9, w - 0.7, h - 1.1, 11)


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float) -> None:
    table = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(w), Cm(h)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(10)
            p.font.bold = r == 0
            p.font.color.rgb = WHITE if r == 0 else BLACK
            p.alignment = PP_ALIGN.CENTER


def add_flow(slide, steps: list[str], x: float, y: float, w: float, h: float) -> None:
    gap = 0.25
    bw = (w - gap * (len(steps) - 1)) / len(steps)
    for i, step in enumerate(steps):
        bx = x + i * (bw + gap)
        color = TEAL if i in {3, 4} else NAVY
        shape = slide.shapes.add_shape(1, Cm(bx), Cm(y), Cm(bw), Cm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color
        p = shape.text_frame.paragraphs[0]
        p.text = step
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER


def add_picture(slide, path: Path, x: float, y: float, w: float, h: float | None = None) -> None:
    if path.exists():
        slide.shapes.add_picture(str(path), Cm(x), Cm(y), width=Cm(w), height=Cm(h) if h else None)
    else:
        add_panel(slide, x, y, w, h or 3.5, "图片缺失", [str(path)])


def main() -> None:
    metrics = load_json(TRAIN_DIR / "metrics.json")
    best = load_json(OPT_DIR / "best_params.json")
    report = load_json(VAL_DIR / "final_validation_report.json")
    sens = load_json(SENS_DIR / "sensitivity_summary.json")
    kfold = load_json(KFOLD_DIR / "kfold_summary.json")
    base3d, opt3d = load_3d_rows()

    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)

    slide = new_slide(prs, "基于 MLP-CNN 流热场重建代理模型的翼型柱阵列散热器优化研究")
    add_bullets(slide, [
        "技术路线：七参数建模 -> CFD 流热场数据 -> MLP-CNN 重建 -> DE 优化 eta -> CFD 复算 -> 三维迁移验证。",
        "核心边界：CFD 云图/数值场不是模型输入，而是监督标签；CNN 用于从参数潜在特征解码空间场。",
    ], 2.0, 5.0, 29.0, 4.0, 18)

    slide = new_slide(prs, "研究问题与目标")
    add_panel(slide, 1.4, 3.0, 9.8, 10.5, "工程问题", ["换热增强和流阻控制需要同时考虑。", "直接用 CFD 搜索高 eta 结构计算成本高。"])
    add_panel(slide, 12.0, 3.0, 9.8, 10.5, "方法问题", ["只预测 Nu/f 可解释性不足。", "重建速度、压力、温度场后，可以解释性能变化来源。"])
    add_panel(slide, 22.6, 3.0, 9.8, 10.5, "本文目标", ["输入七参数，预测 u/v/p/T 场和 Nu/f。", "eta 由 Nu/f 计算，并通过 DE 寻优。", "最终用 CFD 复算和三维模型验证。"])

    slide = new_slide(prs, "完整技术路线")
    add_flow(slide, ["七参数建模", "二维 CFD", "场数据网格化", "MLP-CNN 重建", "DE 优化 eta", "CFD 复算", "三维验证"], 1.2, 3.8, 31.4, 2.2)
    add_bullets(slide, [
        "数据链：case_id、七参数、Nu、f、eta、u/v/p/T 四通道场，统一尺寸 4×50×80。",
        "模型链：MLP Encoder 编码七参数，CNN Decoder 输出流热场，MLP Head 输出 Nu/f。",
        "验证链：代理模型只负责筛选，最终结论以 CFD 复算和三维迁移验证为证据。",
    ], 2.0, 8.0, 29.0, 5.2, 15)

    slide = new_slide(prs, "数据集与损失函数")
    add_table(slide, [
        ["项目", "当前结果"],
        ["样本数", str(metrics["sample_count"])],
        ["训练/验证/测试", f"{metrics['train_count']} / {metrics['val_count']} / {metrics['test_count']}"],
        ["场数据尺寸", "4×50×80"],
        ["场通道", "u、v、p、T"],
        ["损失函数", "L_u+L_v+L_p+L_T+αL_Nu+βL_f+γL_eta"],
        ["权重", f"α={metrics['alpha']}, β={metrics['beta']}, γ={metrics['gamma']}"],
    ], 1.4, 2.7, 16.0, 7.6)
    add_panel(slide, 18.4, 2.7, 13.0, 7.6, "eta 定义", [
        "eta = (Nu/Nu0)/(f/f0)^(1/3)。",
        "Nu0 和 f0 来自基准结构。",
        "eta 同时考虑换热提升和流阻代价。",
    ])

    slide = new_slide(prs, "MLP-CNN 模型结构")
    add_flow(slide, ["七参数", "MLP Encoder", "潜在特征 z", "CNN Decoder", "u/v/p/T"], 1.5, 3.3, 23.5, 2.0)
    add_flow(slide, ["潜在特征 z", "MLP Head", "Nu/f", "eta 计算"], 8.0, 7.0, 17.0, 1.8)
    add_panel(slide, 1.5, 10.2, 30.4, 4.8, "答辩表述", [
        "CNN 不是为了读入云图，而是为了生成空间场。",
        "CFD 数值场作为监督标签，使潜在特征学习流动、压降和温度分布。",
        "性能分支与场重建分支共享参数编码特征，避免纯黑箱标量回归。",
    ])

    slide = new_slide(prs, "训练结果")
    add_table(slide, [
        ["指标", "测试集结果"],
        ["Nu R2", fnum(metrics["Nu_R2"])],
        ["Nu MAE", fnum(metrics["Nu_MAE"])],
        ["f R2", fnum(metrics["f_R2"])],
        ["f MAE", fnum(metrics["f_MAE"])],
        ["eta R2", fnum(metrics["eta_R2"])],
        ["eta MAE", fnum(metrics["eta_MAE"])],
        ["field MAE", fnum(metrics["field_MAE"])],
        ["T MAE", f"{fnum(metrics['T_MAE'])} K"],
    ], 1.4, 2.5, 10.8, 9.8)
    add_picture(slide, TRAIN_DIR / "prediction_scatter.png", 13.1, 2.8, 18.0, 9.8)

    slide = new_slide(prs, "流热场重建效果")
    add_picture(slide, TRAIN_DIR / "case_20040_field_compare.png", 1.0, 2.3, 31.0, 12.6)

    pred = best["prediction"]
    params = best["best_params"]
    slide = new_slide(prs, "DE 优化结果")
    add_table(slide, [["参数", "最优值"]] + [[k, fnum(v)] for k, v in params.items()], 1.4, 2.5, 10.0, 8.6)
    add_panel(slide, 13.0, 2.5, 18.5, 8.6, "代理模型预测", [
        f"Nu_pred = {fnum(pred['Nu_pred'])}",
        f"f_pred = {fnum(pred['f_pred'])}",
        f"eta_pred = {fnum(pred['eta_pred'])}",
        "该候选必须进入 CFD 复算，不能只用代理结果下结论。",
    ])

    proxy = report["proxy_prediction"]
    cfd = report["comsol_result"]
    slide = new_slide(prs, "二维 CFD 复算验证")
    add_table(slide, [
        ["指标", "MLP-CNN 预测", "CFD 复算", "误差"],
        ["Nu", fnum(proxy["Nu_pred"]), fnum(cfd["Nu"]), fnum(proxy["Nu_pred"] - cfd["Nu"])],
        ["f", fnum(proxy["f_pred"]), fnum(cfd["f"]), fnum(proxy["f_pred"] - cfd["f"])],
        ["eta", fnum(proxy["eta_pred"]), fnum(cfd["eta_cfd"]), fnum(proxy["eta_pred"] - cfd["eta_cfd"])],
    ], 1.4, 2.5, 18.6, 4.8)
    add_panel(slide, 21.0, 2.5, 10.8, 4.8, "复算结论", [
        f"CFD eta = {fnum(cfd['eta_cfd'])}。",
        "高于旧 1000 组 baseline ANN/MLP 复核 eta = 1.420111。",
        "代理模型用于筛选，最终数值以 CFD 复算为准。",
    ])
    add_picture(slide, VAL_DIR / "comsol" / "case_1" / "cfd_solution" / "temperature.png", 1.6, 8.5, 9.8, 6.0)
    add_picture(slide, VAL_DIR / "comsol" / "case_1" / "cfd_solution" / "velocity_magnitude.png", 12.0, 8.5, 9.8, 6.0)
    add_picture(slide, VAL_DIR / "comsol" / "case_1" / "cfd_solution" / "pressure.png", 22.4, 8.5, 9.8, 6.0)

    slide = new_slide(prs, "三维迁移验证")
    add_table(slide, [
        ["指标", "三维 baseline", "优化结构", "变化"],
        ["Tavg K", fnum(base3d["t_chip_avg_k"]), fnum(opt3d["t_chip_avg_k"]), f"{float(opt3d['t_chip_avg_k']) - float(base3d['t_chip_avg_k']):+.6f} K"],
        ["Tmax K", fnum(base3d["t_chip_max_k"]), fnum(opt3d["t_chip_max_k"]), f"{float(opt3d['t_chip_max_k']) - float(base3d['t_chip_max_k']):+.6f} K"],
        ["Rth K/W", fnum(base3d["r_th_k_per_w"]), fnum(opt3d["r_th_k_per_w"]), pct_change(opt3d["r_th_k_per_w"], base3d["r_th_k_per_w"])],
        ["delta_p Pa", fnum(base3d["delta_p_pa"]), fnum(opt3d["delta_p_pa"]), pct_change(opt3d["delta_p_pa"], base3d["delta_p_pa"])],
        ["eta3D", fnum(base3d["eta_3d"]), fnum(opt3d["eta_3d"]), "正收益"],
    ], 1.4, 2.5, 18.8, 6.8)
    add_panel(slide, 21.2, 2.5, 10.6, 6.8, "表述边界", [
        "三维是迁移验证，不是三维全局优化。",
        "新版候选在三维中同时降低芯片温度和压降。",
        f"eta3D = {fnum(opt3d['eta_3d'])}。",
    ])
    add_picture(slide, IMG_3D / f"{CASE_3D}_temperature.png", 1.6, 10.0, 9.8, 5.6)
    add_picture(slide, IMG_3D / f"{CASE_3D}_velocity.png", 12.0, 10.0, 9.8, 5.6)
    add_picture(slide, IMG_3D / f"{CASE_3D}_pressure.png", 22.4, 10.0, 9.8, 5.6)

    slide = new_slide(prs, "参数敏感性分析", "基于1000组CFD标签和当前MLP-CNN代理模型")
    add_picture(slide, SENS_DIR / "sensitivity_bar_eta.png", 1.3, 2.6, 9.8, 5.8)
    add_picture(slide, SENS_DIR / "sensitivity_bar_Nu_f.png", 11.8, 2.6, 9.8, 5.8)
    add_picture(slide, SENS_DIR / "local_perturbation_eta.png", 22.3, 2.6, 9.8, 5.8)
    top_rf = sens["eta_top_rf_permutation"][0]
    top_local = sens["eta_top_local"][0]
    add_panel(slide, 1.4, 9.7, 30.5, 4.2, "主要结论", [
        f"eta随机森林置换重要性最高参数：{top_rf['parameter']}，importance={top_rf['rf_permutation_importance']:.6f}。",
        f"DE最优点附近局部扰动最敏感参数：{top_local['parameter']}，eta_range={top_local['eta_range']:.6f}。",
        "该结果是统计敏感性和代理局部扰动分析，不等同于严格CFD单因素扫描。",
    ])

    slide = new_slide(prs, "K折短训稳定性验证", "当前MLP-CNN主线的3折补充验证")
    ks = kfold["summary"]
    add_table(slide, [
        ["指标", "均值", "标准差"],
        ["Nu_R2", fnum(ks["Nu_R2_mean"]), fnum(ks["Nu_R2_std"])],
        ["Nu_MAE", fnum(ks["Nu_MAE_mean"]), fnum(ks["Nu_MAE_std"])],
        ["f_R2", fnum(ks["f_R2_mean"]), fnum(ks["f_R2_std"])],
        ["f_MAE", fnum(ks["f_MAE_mean"]), fnum(ks["f_MAE_std"])],
        ["eta_R2", fnum(ks["eta_R2_mean"]), fnum(ks["eta_R2_std"])],
        ["eta_MAE", fnum(ks["eta_MAE_mean"]), fnum(ks["eta_MAE_std"])],
        ["field_RMSE", fnum(ks["field_RMSE_mean"]), fnum(ks["field_RMSE_std"])],
        ["T_MAE", fnum(ks["T_MAE_mean"]), fnum(ks["T_MAE_std"])],
    ], 1.4, 2.5, 17.0, 9.6)
    add_panel(slide, 19.5, 2.5, 12.0, 9.6, "讲法边界", [
        "3折验证用于说明不同划分下模型可稳定收敛。",
        "每折只训练5 epoch，因此不替代10 epoch正式模型指标。",
        "正式精度仍以mlp_cnn_field_multiloss_e10测试集结果为准。",
    ])

    slide = new_slide(prs, "物理约束与注意力机制", "能证明的写成结果，未训练的写成扩展设计")
    add_panel(slide, 1.4, 2.7, 14.8, 9.8, "当前已实证使用", [
        "L_eta指标一致性约束：eta由Nu/f计算，不作为独立黑箱输出。",
        "u/v/p/T场监督：用CFD数值场约束潜在特征。",
        "损失函数：L_u+L_v+L_p+L_T+alpha L_Nu+beta L_f+gamma L_eta。",
    ])
    add_panel(slide, 17.2, 2.7, 14.8, 9.8, "扩展与边界", [
        "PDE残差、边界条件损失目前是可扩展设计，不写成已完成PINN结果。",
        "旧CBAM/PI-CNN-CBAM属于几何掩码或图像输入路线，只能作历史对照。",
        "当前MLP-CNN注意力版尚未独立训练，不能虚构消融数值。",
    ])

    slide = new_slide(prs, "消融、K 折与论文表达")
    add_panel(slide, 1.4, 2.8, 14.7, 9.8, "消融怎么写", [
        "旧 ANN/MLP 作为 baseline，对比纯标量回归。",
        "MLP-CNN 增加 u/v/p/T 重建监督，说明 CNN 的空间场建模作用。",
        "500 组与 1000 组作为数据规模对照，说明样本量对稳定性的影响。",
    ])
    add_panel(slide, 17.2, 2.8, 14.7, 9.8, "K 折怎么写", [
        "K 折用于证明一次随机划分不是偶然结果。",
        "报告 Nu、f、eta 的均值和标准差。",
        "场重建报告 field_RMSE、T_MAE、p_MAE 的均值和标准差。",
    ])

    slide = new_slide(prs, "老师可能追问")
    add_bullets(slide, [
        "问：这是不是几何掩码？答：最终主线不是几何掩码，而是参数到流热场的 MLP-CNN 重建。",
        "问：CNN 学到了什么？答：CNN 解码 u/v/p/T 空间场，学习局部空间分布和热流耦合结构。",
        "问：为什么还要 CFD 复算？答：代理模型用于快速筛选，最终性能必须由 CFD 复算确认。",
        "问：三维是不是全局最优？答：不是。三维只做迁移验证，结论是二维优化候选在三维模型中仍有效。",
        "问：eta 是什么？答：eta = (Nu/Nu0)/(f/f0)^(1/3)，同时考虑换热增强和流阻代价。",
    ], 1.6, 2.7, 30.5, 11.2, 15)

    slide = new_slide(prs, "结论")
    add_bullets(slide, [
        f"建立了七参数到 u/v/p/T 流热场及 Nu/f 的 MLP-CNN 代理模型，测试集 eta R2 = {fnum(metrics['eta_R2'])}。",
        f"DE 得到候选结构，二维 CFD 复算 eta_cfd = {fnum(cfd['eta_cfd'])}。",
        f"三维芯片级模型迁移验证 eta3D = {fnum(opt3d['eta_3d'])}，当前表现为正收益。",
        "论文主线统一为：参数化建模 + CFD 数据集 + MLP-CNN 流热场重建 + DE 优化 + CFD 复算 + 三维迁移验证。",
    ], 2.0, 3.2, 29.0, 10.0, 17)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
